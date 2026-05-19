#!/usr/bin/env python3
"""Read data/dashboard.xlsx and emit data/data.json for the dashboard.

The Evasion Rate sheet contains three side-by-side tables that this script
splits into separate sections of the JSON output.
"""

from __future__ import annotations

import json
import math
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "dashboard.xlsx"
OUT = ROOT / "data" / "data.json"


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and math.isnan(v)) else v
    s = str(v).strip().replace(",", "").replace("$", "").replace("%", "")
    if not s or s == "-":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def txt(v):
    if v is None:
        return ""
    return str(v).strip()


def parse_evasion_rate(ws):
    rows = list(ws.iter_rows(values_only=True))
    # Header row is the first row whose first cell is exactly "Route".
    header_row = next(i for i, r in enumerate(rows) if txt(r[0]) == "Route")
    data_rows = rows[header_row + 1 :]

    bus, rail, totals = [], [], []

    for r in data_rows:
        # Bus routes (cols A-F)
        route, b_apc, b_adj, b_paid, b_rate, b_notes = r[0], r[1], r[2], r[3], r[4], r[5]
        if txt(route) or num(b_apc) is not None:
            bus.append(
                {
                    "route": txt(route),
                    "apc_boardings": num(b_apc),
                    "adjusted_boardings": num(b_adj),
                    "paid_entries": num(b_paid),
                    "evasion_rate": num(b_rate),
                    "notes": txt(b_notes),
                }
            )

        # Rail/G-Line stations (cols G-M)
        merged, r_apc, r_adj, r_station, r_paid, r_rate, r_notes = (
            r[6], r[7], r[8], r[9], r[10], r[11], r[12],
        )
        if txt(merged) or num(r_apc) is not None:
            rail.append(
                {
                    "merged_station": txt(merged),
                    "apc_boardings": num(r_apc),
                    "adjusted_boardings": num(r_adj),
                    "fare_station": txt(r_station),
                    "paid_entries": num(r_paid),
                    "evasion_rate": num(r_rate),
                    "notes": txt(r_notes),
                }
            )

        # Mode totals (cols N-R)
        mode, t_apc, t_paid, t_rate, t_notes = r[13], r[14], r[15], r[16], r[17]
        if txt(mode) or num(t_apc) is not None:
            totals.append(
                {
                    "mode": txt(mode),
                    "apc_boardings": num(t_apc),
                    "paid_boardings": num(t_paid),
                    "evasion_rate": num(t_rate),
                    "notes": txt(t_notes),
                }
            )

    return {"bus": bus, "rail": rail, "totals": totals}


def parse_intro(ws):
    rows = list(ws.iter_rows(values_only=True))
    title = txt(rows[0][0]) if rows else ""
    blurb = txt(rows[1][0]) if len(rows) > 1 else ""
    return {"title": title, "blurb": blurb}


def main():
    if not XLSX.exists():
        print(f"Missing {XLSX}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Evasion Rate"]

    payload = {
        "intro": parse_intro(ws),
        **parse_evasion_rate(ws),
        "sheets_available": wb.sheetnames,
    }

    OUT.write_text(json.dumps(payload, indent=2, default=str))
    print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes)")
    print(
        f"  bus rows: {len(payload['bus'])}, rail rows: {len(payload['rail'])}, totals rows: {len(payload['totals'])}"
    )


if __name__ == "__main__":
    main()
